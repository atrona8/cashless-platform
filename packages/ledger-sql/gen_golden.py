# Scénario de référence exécutable.
#   python3 gen_golden.py                      : lit scenario_reference.json (SOURCE) et génère scenario_reference_test.sql
#   python3 gen_golden.py --from-xlsx FICHIER  : régénère d'abord scenario_reference.json depuis le classeur recalculé
#                                                (outil d'auteur ; le classeur n'est pas normatif)
# Le SQL généré crée le grand livre, rejoue chaque transaction avec post_transaction, puis compare les soldes
# de la base à ceux du scénario (fin du festival et fin de clôture).
import json, sys
LIVE_END = 23                       # dernière transaction « pendant le festival »

TYPE = {  # libellé du classeur -> type de transaction en base
    "Recharge": "TOPUP", "Recharge espèces": "TOPUP_CASH", "Caution": "DEPOSIT_TAKEN", "Caution rendue": "DEPOSIT_REFUNDED",
    "Caution acquise": "DEPOSIT_FORFEITED", "Crédits offerts": "PROMO_CREDIT", "Frais d'activation": "ACTIVATION_FEE",
    "Vente": "PURCHASE", "Annulation vente": "REVERSAL", "Vente hors ligne": "PURCHASE", "Fermeture caisse": "CASH_CLOSE",
    "Résolution anomalie": "ANOMALY_RESOLUTION", "Dépôt espèces": "CASH_DEPOSIT", "Versement PSP": "PSP_SETTLEMENT",
    "Droits de place": "PITCH_FEE", "Frais prestataire": "OPERATOR_FEE", "Redevance plateforme": "PLATFORM_FEE",
    "Versements initiés": "PAYOUT_INITIATED", "Versement initié": "PAYOUT_INITIATED", "Versement confirmé": "PAYOUT_CONFIRMED",
    "Versement échoué": "PAYOUT_FAILED", "Remboursement": "WALLET_REFUND", "Casse": "BREAKAGE",
    "Expiration crédits offerts": "PROMO_EXPIRY",
}
FAMILY = {"Argent": "ASSET", "Droit": "CLAIM", "Attente": "SUSPENSE"}
PARTIES = {  # titulaire du classeur -> (uuid, kind, nom)
    "PLT": ("00000000-0000-0000-0000-00000000a001", "PLATFORM", "Plateforme Cashless"),
    "OPE": ("00000000-0000-0000-0000-00000000a002", "OPERATOR", "Prestataire Sénégal"),
    "ORG": ("00000000-0000-0000-0000-00000000a003", "ORGANIZER", "Association Festival Sons"),
    "BAR": ("00000000-0000-0000-0000-00000000a004", "MERCHANT", "Bar de l'organisateur"),
    "FOOD": ("00000000-0000-0000-0000-00000000a005", "MERCHANT", "Food truck"),
    "TEE": ("00000000-0000-0000-0000-00000000a006", "MERCHANT", "Vendeur de tee-shirts"),
}
LEDGER = "10000000-0000-0000-0000-00000000a001"
USER1, USER2 = "90000000-0000-0000-0000-00000000a001", "90000000-0000-0000-0000-00000000a002"

def extract_from_xlsx(XLSX):
    from openpyxl import load_workbook
    wb = load_workbook(XLSX, data_only=True)
    accounts = []
    for row in wb["Plan de comptes"].iter_rows(min_row=5, values_only=True):
        if not row[0]: break
        code, label, fam, owner, side, neg, purpose = row[:7]
        accounts.append({"code": code, "label": label, "family": FAMILY[fam], "purpose": purpose,
                         "owner": owner if owner in PARTIES else None, "normal_side": side, "allow_negative": neg == "Oui"})

    txs = {}
    for row in wb["Journal"].iter_rows(min_row=5, values_only=True):
        if not isinstance(row[0], int): continue
        n, date, typ, nat, lib, code, _, deb, cred = row[:9]
        src, key = row[11], row[12]
        t = txs.setdefault(n, {"no": n, "occurred_at": date.strftime("%Y-%m-%dT%H:%M:00+00:00"), "label": typ,
                               "type": TYPE[typ], "source": src, "idempotency_key": key, "lines": []})
        amount = int(round((deb or 0) - (cred or 0)))
        if amount:                                   # une ligne à 0 (ex. écart de caisse nul d'un côté) n'est pas écrite
            t["lines"].append({"account": code, "amount": amount, "memo": lib})
    transactions = [txs[k] for k in sorted(txs)]
    for i, t in enumerate(transactions):
        if t["type"] == "REVERSAL":           # dans le scénario, l'annulation suit immédiatement la vente annulée
            t["reverses_key"] = transactions[i - 1]["idempotency_key"]
    for t in transactions:
        assert sum(l["amount"] for l in t["lines"]) == 0, t

    def balances_after(n):
        b = {a["code"]: 0 for a in accounts}
        for t in transactions:
            if t["no"] <= n:
                for l in t["lines"]: b[l["account"]] += l["amount"]
        return b
    # contrôle croisé : les soldes finaux recalculés ici = la feuille Balance d'Excel (formules SUMIFS)
    xl_final = {r[0]: r[6] for r in wb["Balance"].iter_rows(min_row=5, values_only=True) if r[0] and r[0] in {a["code"] for a in accounts}}
    final = balances_after(10**6)
    assert all(final[c] == xl_final[c] for c in final), "écart entre le journal extrait et la Balance du classeur"

    ref = {
        "name": "Festival Sons 2026 — scénario de référence",
        "status": "NORMATIF pour les tests : tout changement du moteur d'écritures doit laisser ce scénario vert, ou le modifier explicitement.",
        "currency": "XOF", "minor_units": 0,
        "parameters": {r[0]: r[2] for r in wb["Paramètres"].iter_rows(min_row=5, values_only=True) if r[0]},
        "parties": {k: {"id": v[0], "kind": v[1], "name": v[2]} for k, v in PARTIES.items()},
        "accounts": accounts, "transactions": transactions,
        "expected": {
            "after_festival": {"after_tx": LIVE_END, "balances": balances_after(LIVE_END)},
            "after_closing": {"after_tx": transactions[-1]["no"], "balances": final},
        },
    }
    json.dump(ref, open("scenario_reference.json", "w"), ensure_ascii=False, indent=1)


if "--from-xlsx" in sys.argv:
    extract_from_xlsx(sys.argv[sys.argv.index("--from-xlsx") + 1])
ref = json.load(open("scenario_reference.json"))
accounts, transactions = ref["accounts"], ref["transactions"]
final = ref["expected"]["after_closing"]["balances"]
LIVE_END = ref["expected"]["after_festival"]["after_tx"]

# ---------------------------------------------------------------- pgTAP
def q(x): return "$Q$" + x + "$Q$"
def lit(x): return "'" + x.replace("'", "''") + "'"
out = []; n_tests = 0
def t(sql):
    global n_tests; n_tests += 1; out.append(sql)
op = PARTIES["OPE"][0]
out.append(f"SET app.operator_id = '{op}';")
for k, (pid, kind, name) in PARTIES.items():
    opcol = "NULL" if kind in ("PLATFORM", "OPERATOR") else f"'{op}'"
    out.append(f"INSERT INTO party (id, kind, operator_id, legal_name, country_code) VALUES ('{pid}','{kind}',{opcol},{q(name)},'SN');")
org = PARTIES["ORG"][0]
EVENT, JUR = "60000000-0000-0000-0000-00000000a001", "50000000-0000-0000-0000-00000000a001"
out.append(f"INSERT INTO jurisdiction_profile (id, country_code, version, valid_from, breakage_destination) "
           f"VALUES ('{JUR}','SN',901,'2026-01-01','ORGANIZER');")
out.append(f"INSERT INTO event (id, operator_id, organizer_id, name, currency, timezone, jurisdiction_id, funds_holder) "
           f"VALUES ('{EVENT}','{op}','{org}',{q(ref['name'])},'XOF','Africa/Dakar','{JUR}','ORGANIZER');")
out.append(f"INSERT INTO ledger (id, operator_id, scope_type, scope_id, currency, issuer_id, funds_holder_id) "
           f"VALUES ('{LEDGER}','{op}','EVENT','{EVENT}','XOF','{org}','{org}');")
for a in accounts:
    owner = f"'{PARTIES[a['owner']][0]}'" if a["owner"] else "NULL"
    out.append(f"INSERT INTO account (operator_id, ledger_id, code, family, purpose, owner_party_id, normal_side, allow_negative) "
               f"VALUES ('{op}','{LEDGER}','{a['code']}','{a['family']}','{a['purpose']}',{owner},'{a['normal_side']}',{str(a['allow_negative']).lower()});")
# comptes « chauds » comme en production : commissions, taxes, créances PSP
out.append("UPDATE account SET hot = true WHERE code IN ('L-ORG-COM','L-ORG-TVA','A-PSP-WAVE','A-PSP-OM','A-PSP-CARTE');")

def post_sql(tx):
    lines = json.dumps([{"account_id": "@" + l["account"], "amount": l["amount"], "memo": l["memo"]} for l in tx["lines"]], ensure_ascii=False)
    # les codes sont remplacés par les identifiants au moment de l'appel
    lines_expr = (f"(SELECT jsonb_agg(jsonb_set(e, '{{account_id}}', to_jsonb((SELECT id::text FROM account WHERE ledger_id = '{LEDGER}' "
                  f"AND code = substr(e->>'account_id', 2))))) FROM jsonb_array_elements({lit(lines)}::jsonb) e)")
    rev = (f"(SELECT id FROM journal_transaction WHERE ledger_id = '{LEDGER}' AND idempotency_key = {lit(tx['reverses_key'])})"
           if tx.get("reverses_key") else "NULL")
    appr = f", NULL, {rev}, '{USER1}', '{USER2}'" if tx["source"] == "BACKOFFICE" else (f", NULL, {rev}" if tx.get("reverses_key") else "")
    return (f"SELECT post_transaction('{LEDGER}', '{tx['type']}', {lit(tx['idempotency_key'])}, '{tx['occurred_at']}', "
            f"'{tx['source']}', {lines_expr}{appr})")

def check_balances(label, bal, at=None):
    src = (f"trial_balance_at('{LEDGER}', '{at}')" if at else
           f"(SELECT code, signed_balance FROM trial_balance WHERE ledger_id = '{LEDGER}') tb")
    exp = ", ".join(f"('{c}', {v})" for c, v in bal.items())
    t(f"SELECT results_eq({q(f'SELECT code, signed_balance::bigint FROM {src} ORDER BY code')}, "
      f"{q(f'SELECT code, bal::bigint FROM (VALUES {exp}) v(code, bal) ORDER BY code')}, {q(label)});")

for tx in transactions:
    label = "T%02d %s (%s)" % (tx["no"], tx["label"], tx["type"])
    t(f"SELECT lives_ok({q(post_sql(tx))}, {q(label)});")
live_at = transactions[LIVE_END - 1]["occurred_at"]
check_balances(f"fin du festival (après T{LIVE_END}) : soldes = scénario", ref["expected"]["after_festival"]["balances"], at=live_at)
check_balances("fin de clôture : soldes = scénario", final)
t(f"SELECT is((SELECT must_be_zero FROM ledger_invariant WHERE ledger_id = '{LEDGER}')::text, '0', 'invariant : argent = droits');")
t(f"SELECT is((SELECT count(*) FROM party_position WHERE ledger_id = '{LEDGER}' AND net_claim <> 0)::int, 0, 'toutes les parties sont soldées');")
t(f"SELECT is((SELECT count(*) FROM balance_drift)::int, 0, 'soldes en cache = recalcul');")
t(f"SELECT is((SELECT count(*) FROM hot_account_side_check)::int, 0, 'comptes chauds du bon côté');")
# rejouer tout le scénario ne crée rien (idempotence de bout en bout)
out.append(f"DO $R$ BEGIN {' '.join('PERFORM ' + post_sql(tx)[7:] + ';' for tx in transactions)} END $R$;")
t(f"SELECT is((SELECT count(*) FROM journal_transaction WHERE ledger_id = '{LEDGER}')::int, {len(transactions)}, 'rejeu complet du scénario : aucun doublon');")

# Clôture complète du scénario (revue 2, DB-B1) : les comptes d'une même partie se compensent (L-ORG-COM,
# L-ORG-VERS...) ; la clôture doit aboutir et verrouiller le grand livre avec son scellement final.
for st in ("LIVE", "CLOSING", "RECONCILING", "SETTLING", "REFUND_WINDOW", "CLOSED"):
    call = f"SELECT set_event_status('{EVENT}', '{st}', '{USER1}')"
    t(f"SELECT lives_ok({q(call)}, {q('statut de l événement : ' + st)});")
t(f"SELECT is((SELECT status FROM ledger WHERE id = '{LEDGER}'), 'LOCKED', 'scénario clos : grand livre verrouillé');")
t(f"SELECT is((SELECT bool_and(ok) FROM verify_ledger_seals('{LEDGER}')), true, 'scellement final vérifié');")
t(f"SELECT is((SELECT remaining FROM event_status_history WHERE event_id = '{EVENT}' AND to_status = 'CLOSED'), 0::bigint, 'aucun solde restant à la clôture');")

hdr = ["-- Scénario de référence rejoué contre la base (pgTAP). Généré par gen_golden.py depuis scenario_reference.json.",
       "-- Lancer : pg_prove -d <base> scenario_reference_test.sql",
       "BEGIN;", "CREATE EXTENSION IF NOT EXISTS pgtap;", f"SELECT plan({n_tests});"]
open("scenario_reference_test.sql", "w").write("\n".join(hdr + out + ["SELECT * FROM finish();", "ROLLBACK;"]) + "\n")
print("transactions", len(transactions), "tests", n_tests)
